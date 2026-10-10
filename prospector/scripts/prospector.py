#!/usr/bin/env python3
"""
Prospector v4 - Buscador general de negocios en Google Maps (Apify) -> XLSX + CSV

Generico: sirve para cualquier nicho y cualquier zona. Acepta VARIOS terminos de
busqueda y VARIAS zonas en una sola corrida, y consolida todo en un unico archivo
con deduplicacion global.

Si faltan datos y hay terminal, los pregunta. Si no hay terminal, exige los flags.

    export APIFY_TOKEN=...
    python3 prospector.py                          # modo interactivo
    python3 prospector.py --receta rx-medico --zona "San Jose, Costa Rica"
    python3 prospector.py --buscar "bufete de abogados" "notaria" --zona Heredia Cartago
"""

import argparse
import csv
import json
import os
import re
import sys
import time
from datetime import date
from urllib.parse import quote

import requests

ACTOR_ID = "nwua9Gu5YrADL7ZDj"  # compass/crawler-google-places
APIFY = "https://api.apify.com/v2"

# Token fuera del codigo (2026-10-10): se lee de la variable de entorno de usuario APIFY_TOKEN.
# El flag --apify-key tiene precedencia. Motivo: GitHub bloquea secretos en el repositorio.
APIFY_TOKEN_EMBEBIDO = os.environ.get("APIFY_TOKEN", "")

# Costo observado 2026-07-28: USD 0.1252 / 25 lugares = USD 0.0050 por lugar.
USD_POR_LUGAR = 0.005

CALLING_CODES = {
    "CR": "506", "PA": "507", "NI": "505", "HN": "504", "SV": "503", "GT": "502",
    "BZ": "501", "MX": "52", "US": "1", "CA": "1", "DO": "1", "PR": "1",
    "CO": "57", "VE": "58", "EC": "593", "PE": "51", "BO": "591", "CL": "56",
    "AR": "54", "PY": "595", "UY": "598", "BR": "55",
    "ES": "34", "PT": "351", "FR": "33", "IT": "39", "DE": "49", "GB": "44",
}

# Recetas: paquetes de terminos de busqueda para casos frecuentes.
# Google Maps indexa por CATEGORIA, no por capacidad instalada: "centros medicos
# con rayos X" no es una categoria. Por eso cada receta barre varios terminos.
RECETAS = {
    "rx-medico": {
        "terminos": ["centro de radiologia", "centro de imagenes medicas",
                     "rayos x", "clinica medica", "hospital privado"],
        "nota": "Compradores de equipo e insumos de proteccion radiologica medica. "
                "El resultado NO confirma equipo instalado: calificar contra el "
                "registro de instalaciones del MINSA (Decreto 44653-S).",
    },
    "dental": {
        "terminos": ["clinica dental", "odontologo", "ortodoncista"],
        "nota": "RX dental, panoramico y CBCT. Alta densidad; suele requerir "
                "--max-results alto o barrido por canton.",
    },
    "veterinaria": {
        "terminos": ["clinica veterinaria", "veterinario", "hospital veterinario"],
        "nota": "Nicho menos atendido en RX. Fichas con muchas resenas suelen ser "
                "cadenas con una sola decision de compra.",
    },
    "industrial": {
        "terminos": ["ensayos no destructivos", "laboratorio de materiales",
                     "inspeccion industrial", "soldadura industrial"],
        "nota": "Gammagrafia industrial y radiografia de soldadura. Volumen bajo, "
                "ticket alto.",
    },
    "inmobiliaria": {
        "terminos": ["inmobiliaria", "bienes raices", "corredor de bienes raices",
                     "administracion de condominios"],
        "nota": "Prospeccion patrimonial personal. NO aplicar identidad Cole S.A.",
    },
    "hospedaje": {
        "terminos": ["hotel", "apartotel", "hostal", "cabinas"],
        "nota": "Prospeccion patrimonial personal. NO aplicar identidad Cole S.A.",
    },
}

ESTADOS = ["Nuevo", "Contactado", "Interesado", "Cotizado", "Ganado", "Descartado"]
COLUMNAS = ["Empresa", "Telefono", "WhatsApp", "Rating", "Resenas", "Categoria",
            "Direccion", "Ciudad", "Zona buscada", "Web", "Google Maps",
            "Estado", "Capturado", "Notas"]

AZUL, GRIS = "3155E2", "363B47"
FUENTE = "Arial"


# ------------------------------------------------------------------ HTTP

def request(method, url, *, retries=4, **kw):
    delay, last = 2, None
    for intento in range(retries + 1):
        try:
            r = requests.request(method, url, **kw)
        except requests.RequestException as e:
            last = e
            if intento == retries:
                raise
            time.sleep(delay); delay *= 2
            continue
        if r.status_code == 429 or 500 <= r.status_code < 600:
            if intento == retries:
                return r
            time.sleep(min(float(r.headers.get("Retry-After", delay)), 30))
            delay *= 2
            continue
        return r
    raise last  # pragma: no cover


# ------------------------------------------------------------------ Interactivo

def hay_terminal():
    return sys.stdin.isatty() and sys.stdout.isatty()


def preguntar(texto, default=""):
    if not hay_terminal():
        return default
    sufijo = f" [{default}]" if default else ""
    try:
        r = input(f"  ? {texto}{sufijo}: ").strip()
    except (EOFError, KeyboardInterrupt):
        print(); sys.exit("[X] Cancelado.")
    return r or default


def intake(args):
    """Completa lo que falte preguntando. Sin terminal, exige los flags."""
    if args.receta:
        if args.receta not in RECETAS:
            sys.exit(f"[X] Receta desconocida: {args.receta}. "
                     f"Disponibles: {', '.join(RECETAS)}")
        args.niche = args.niche or RECETAS[args.receta]["terminos"]

    if not args.niche:
        if not hay_terminal():
            sys.exit("[X] Falta --buscar/--niche (o --receta). "
                     f"Recetas: {', '.join(RECETAS)}")
        print("\n  Recetas disponibles:")
        for k, v in RECETAS.items():
            print(f"    {k:14} {', '.join(v['terminos'][:3])}...")
        r = preguntar("Receta a usar (Enter para escribir terminos libres)")
        if r and r in RECETAS:
            args.receta = r
            args.niche = RECETAS[r]["terminos"]
        else:
            libre = preguntar("Que buscas (separa varios terminos con ';')")
            if not libre:
                sys.exit("[X] Sin termino de busqueda no hay corrida.")
            args.niche = [t.strip() for t in libre.split(";") if t.strip()]

    if not args.zone:
        z = preguntar("Donde (separa varias zonas con ';')", "San Jose, Costa Rica")
        args.zone = [t.strip() for t in z.split(";") if t.strip()]

    if not args.country and hay_terminal():
        args.country = preguntar("Codigo ISO2 del pais (Enter = geocodificar)", "") or None

    if hay_terminal() and args.max_results == 200:
        m = preguntar("Tope de lugares POR TERMINO Y POR ZONA", "200")
        try:
            args.max_results = max(1, int(m))
        except ValueError:
            pass

    return args


# ------------------------------------------------------------------ Geocoding

def resolver_pais(zone, override=None):
    if override:
        return override.strip().upper()
    try:
        r = requests.get("https://nominatim.openstreetmap.org/search",
                         params={"q": zone, "format": "json", "limit": 1,
                                 "addressdetails": 1},
                         headers={"User-Agent": "prospector-skill/4.0"}, timeout=10)
        r.raise_for_status()
        res = r.json()
        if res:
            return (res[0].get("address", {}).get("country_code") or "").upper()
    except Exception as e:
        print(f"   [!] Geocoding fallo ({e}); usa --country si el filtro importa")
    return ""


# ------------------------------------------------------------------ Presupuesto

def presupuesto(token, estimado):
    r = request("GET", f"{APIFY}/users/me/limits", params={"token": token}, timeout=20)
    if r.status_code != 200:
        print("      [!] No se pudo leer el credito; se continua sin verificar")
        return
    d = r.json()["data"]
    tope = d["limits"].get("maxMonthlyUsageUsd")
    usado = d["current"].get("monthlyUsageUsd", 0.0)
    fin = d.get("monthlyUsageCycle", {}).get("endAt", "")[:10]
    if tope is None:
        print(f"      [!] La cuenta no declara tope mensual (maxMonthlyUsageUsd). "
              f"La guarda de presupuesto NO se aplica: la corrida estimada en "
              f"USD {estimado:.2f} se ejecuta sin verificar contra credito. "
              f"Uso del ciclo hasta ahora: USD {usado:.2f}.")
        return
    resta = tope - usado
    print(f"      Credito Apify: USD {resta:.2f} de {tope:.2f} (ciclo cierra {fin})")
    if estimado > resta:
        sys.exit(f"[X] Corrida estimada USD {estimado:.2f} > credito restante "
                 f"USD {resta:.2f}. Baja --max-results, quita terminos o zonas, "
                 f"o usa --ignorar-presupuesto.")


# ------------------------------------------------------------------ Apify

def costo_run(token, run_id):
    """Costo real facturado por Apify para un run. 0.0 si no se puede leer.

    Se consulta despues de que el run termina: usageTotalUsd es el cargo real,
    no la tarifa de referencia USD_POR_LUGAR (medida una sola vez el 2026-07-28).
    """
    try:
        r = request("GET", f"{APIFY}/actor-runs/{run_id}",
                    params={"token": token}, timeout=15)
        if r.status_code == 200:
            return float(r.json()["data"].get("usageTotalUsd") or 0.0)
    except Exception:
        pass
    return 0.0


def correr_zona(token, zona, args, country):
    payload = {
        "searchStringsArray": list(args.niche),
        "locationQuery": zona,
        "maxCrawledPlacesPerSearch": args.max_results,
        "language": args.language,
        "skipClosedPlaces": bool(args.skip_cerrados),
        "maxImages": 0, "maxReviews": 0,
        "includeHistogram": False, "includeOpeningHours": False,
        "includeWebResults": False,
    }
    if country:
        payload["countryCode"] = country.lower()
    if args.sin_web:
        payload["website"] = "withoutWebsite"
    # min_rating / min_reviews se filtran localmente: el actor espera un enum
    # ("four", "fourAndHalf") y un valor numerico haria fallar el run.

    r = request("POST", f"{APIFY}/acts/{ACTOR_ID}/runs",
                params={"token": token}, json=payload, timeout=30)
    if r.status_code >= 400:
        sys.exit(f"[X] Apify rechazo el run ({r.status_code}): {r.text[:300]}")
    run_id = r.json()["data"]["id"]
    print(f"      [{zona}] run {run_id}")

    while True:
        r = request("GET", f"{APIFY}/actor-runs/{run_id}",
                    params={"token": token}, timeout=15)
        if r.status_code >= 400:
            sys.exit(f"[X] No se pudo consultar el run: {r.text[:200]}")
        d = r.json()["data"]
        estado, n = d["status"], d.get("stats", {}).get("itemCount", 0)
        print(f"      [{zona}] {estado} | lugares: {n}      ", end="\r", flush=True)
        if estado == "SUCCEEDED":
            print(); break
        if estado in ("FAILED", "ABORTED", "TIMED-OUT", "TIMING-OUT"):
            print()
            print(f"      [!] {zona} termino en {estado}; se omite esa zona")
            return [], run_id, costo_run(token, run_id)
        time.sleep(8)

    leads, offset = [], 0
    while True:
        r = request("GET", f"{APIFY}/actor-runs/{run_id}/dataset/items",
                    params={"token": token, "offset": offset, "limit": 1000}, timeout=90)
        if r.status_code >= 400:
            sys.exit(f"[X] No se pudo bajar el dataset: {r.text[:200]}")
        lote = r.json()
        if not lote:
            break
        leads.extend(lote); offset += len(lote)
        if len(lote) < 1000:
            break
    for l in leads:
        l["_zona"] = zona
    return leads, run_id, costo_run(token, run_id)


def scrapear(token, args, country):
    combinaciones = len(args.niche) * len(args.zone)
    estimado = combinaciones * args.max_results * USD_POR_LUGAR
    print(f"[1/4] Apify: {len(args.niche)} termino(s) x {len(args.zone)} zona(s) "
          f"= {combinaciones} busquedas")
    for t in args.niche:
        print(f"        - {t}")
    print(f"      Zonas: {', '.join(args.zone)} | pais={country or 'n/d'}")
    print(f"      Tope {args.max_results}/busqueda -> maximo {combinaciones * args.max_results}"
          f" lugares ~= USD {estimado:.2f}")

    if not args.ignorar_presupuesto:
        presupuesto(token, estimado)

    if args.dry_run:
        print("      --dry-run: no se ejecuta el actor.")
        sys.exit(0)

    crudos, runs = [], []
    for zona in args.zone:
        leads, run_id, costo = correr_zona(token, zona, args, country)
        crudos.extend(leads)
        runs.append((zona, run_id, len(leads), costo))
    real = sum(c for *_, c in runs)
    if real:
        por_lugar = real / len(crudos) if crudos else 0.0
        aviso = ""
        if por_lugar and abs(por_lugar - USD_POR_LUGAR) / USD_POR_LUGAR > 0.20:
            aviso = (f"  [!] Desvio >20% vs. tarifa de referencia "
                     f"USD {USD_POR_LUGAR:.4f}; revisar USD_POR_LUGAR.")
        print(f"      {len(crudos)} lugares descargados | costo real Apify "
              f"USD {real:.4f} (USD {por_lugar:.4f}/lugar) | estimado previo "
              f"USD {estimado:.2f}{aviso}\n")
    else:
        print(f"      {len(crudos)} lugares descargados | Apify no reporto "
              f"usageTotalUsd; se usa la tarifa de referencia\n")
    return crudos, runs


# ------------------------------------------------------------------ Normalizacion

def normalizar_tel(crudo, country):
    """E.164 sin '+' (formato wa.me) o '' si no es utilizable."""
    if not crudo:
        return ""
    s = str(crudo).strip()
    internacional = s.startswith("+") or s.startswith("00")
    digitos = re.sub(r"\D", "", s)
    if s.startswith("00"):
        digitos = digitos[2:]
    if not digitos:
        return ""
    if internacional:
        return digitos if 8 <= len(digitos) <= 15 else ""
    cc = CALLING_CODES.get((country or "").upper())
    if not cc:
        return ""
    nacional = digitos.lstrip("0")
    if not nacional or len(nacional) < 6:
        return ""
    if nacional.startswith(cc) and len(nacional) > len(cc) + 6:
        return nacional
    completo = cc + nacional
    return completo if 8 <= len(completo) <= 15 else ""


def wa_link(tel, empresa, plantilla):
    if not tel:
        return ""
    return f"https://wa.me/{tel}?text=" + quote(
        plantilla.replace("{empresa}", empresa or "ustedes"), safe="")


def url_valida(u):
    return isinstance(u, str) and u.startswith(("http://", "https://"))


def construir(leads, country, args):
    """Aplana, filtra y deduplica GLOBALMENTE (entre terminos y entre zonas)."""
    filas, vistos_id, vistos_tel = [], set(), set()
    desc = {"duplicado": 0, "sin_nombre": 0, "rating": 0, "resenas": 0,
            "cerrado": 0, "pais": 0, "sin_telefono": 0}
    hoy = date.today().isoformat()

    for l in leads:
        nombre = (l.get("title") or "").strip()
        if not nombre:
            desc["sin_nombre"] += 1; continue

        pid = l.get("placeId") or l.get("cid")
        if pid and pid in vistos_id:
            desc["duplicado"] += 1; continue

        if args.skip_cerrados and (l.get("permanentlyClosed") or l.get("temporarilyClosed")):
            desc["cerrado"] += 1; continue

        lead_cc = (l.get("countryCode") or "").upper()
        if country and lead_cc and lead_cc != country:
            desc["pais"] += 1; continue

        try:
            rating = round(float(l.get("totalScore")), 1) if l.get("totalScore") is not None else None
        except (TypeError, ValueError):
            rating = None
        if args.min_rating and (rating is None or rating < args.min_rating):
            desc["rating"] += 1; continue

        try:
            resenas = int(l.get("reviewsCount") or 0)
        except (TypeError, ValueError):
            resenas = 0
        if args.min_reviews and resenas < args.min_reviews:
            desc["resenas"] += 1; continue

        tel = normalizar_tel(l.get("phoneUnformatted") or l.get("phone"), country)
        if args.solo_con_telefono and not tel:
            desc["sin_telefono"] += 1; continue
        if tel and tel in vistos_tel:
            desc["duplicado"] += 1; continue

        cat = l.get("categoryName") or ""
        if not cat and isinstance(l.get("categories"), list) and l["categories"]:
            cat = str(l["categories"][0])

        filas.append({
            "Empresa": nombre,
            "Telefono": ("+" + tel) if tel else "",
            "WhatsApp": wa_link(tel, nombre, args.mensaje),
            "Rating": rating,
            "Resenas": resenas,
            "Categoria": cat,
            "Direccion": l.get("address") or "",
            "Ciudad": l.get("city") or "",
            "Zona buscada": l.get("_zona", ""),
            "Web": l["website"] if url_valida(l.get("website")) else "",
            "Google Maps": l["url"] if url_valida(l.get("url")) else "",
            "Estado": "Nuevo",
            "Capturado": hoy,
            "Notas": "",
        })
        if pid:
            vistos_id.add(pid)
        if tel:
            vistos_tel.add(tel)

    return filas, desc


# ------------------------------------------------------------------ Salidas

def escribir_csv(filas, ruta):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNAS, extrasaction="ignore")
        w.writeheader(); w.writerows(filas)


def escribir_xlsx(filas, ruta, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    anchos = {"Empresa": 34, "Telefono": 15, "WhatsApp": 13, "Rating": 8,
              "Resenas": 9, "Categoria": 22, "Direccion": 42, "Ciudad": 18,
              "Zona buscada": 20, "Web": 13, "Google Maps": 13, "Estado": 14,
              "Capturado": 12, "Notas": 34}
    centradas = {"WhatsApp", "Rating", "Resenas", "Web", "Google Maps",
                 "Estado", "Capturado"}
    enlaces = {"WhatsApp": "Escribir", "Web": "Sitio", "Google Maps": "Ficha"}
    col = {h: i + 1 for i, h in enumerate(COLUMNAS)}

    wb = Workbook(); ws = wb.active; ws.title = "Leads"
    thin = Side(style="thin", color="D9D9D9")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    for h, i in col.items():
        c = ws.cell(1, i, h)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = anchos[h]
    ws.row_dimensions[1].height = 28

    for r, d in enumerate(filas, 2):
        for h, i in col.items():
            c = ws.cell(r, i)
            if h in enlaces:
                if d[h]:
                    c.value = enlaces[h]; c.hyperlink = d[h]
                    c.font = Font(name=FUENTE, size=10, color="0563C1", underline="single")
                else:
                    c.value = "-"; c.font = Font(name=FUENTE, size=10, color="A6A6A6")
            else:
                c.value = d[h]
                c.font = Font(name=FUENTE, size=10)
            c.border = borde
            c.alignment = Alignment(vertical="center",
                                    wrap_text=h in ("Empresa", "Direccion"),
                                    horizontal="center" if h in centradas else "left")
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F4F7FE")
        ws.cell(r, col["Rating"]).number_format = "0.0"
        ws.cell(r, col["Resenas"]).number_format = "#,##0"
        ws.cell(r, col["Capturado"]).number_format = "YYYY-MM-DD"

    n = len(filas) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{n}"
    ws.freeze_panes = "B2"
    letra_estado = get_column_letter(col["Estado"])
    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"',
                        allow_blank=True, showDropDown=False, promptTitle="Pipeline",
                        prompt="Actualiza el estado conforme avance la gestion.")
    ws.add_data_validation(dv)
    dv.add(f"{letra_estado}2:{letra_estado}{n}")

    L = {h: get_column_letter(i) for h, i in col.items()}

    # --- Resumen (todo por formula)
    rs = wb.create_sheet("Resumen")
    rs.column_dimensions["A"].width = 30
    rs.column_dimensions["B"].width = 14
    rs.column_dimensions["C"].width = 12
    rs.cell(1, 1, "Resumen del barrido").font = Font(name=FUENTE, bold=True,
                                                     size=14, color=GRIS)
    metricas = [
        ("Leads en la hoja", f"=COUNTA(Leads!{L['Empresa']}2:{L['Empresa']}{n})", None),
        ("Con telefono utilizable",
         f'=COUNTIF(Leads!{L["WhatsApp"]}2:{L["WhatsApp"]}{n},"Escribir")', None),
        ("Con sitio web", f'=COUNTIF(Leads!{L["Web"]}2:{L["Web"]}{n},"Sitio")', None),
        ("Sin sitio web", f'=COUNTIF(Leads!{L["Web"]}2:{L["Web"]}{n},"-")', None),
        ("Rating promedio",
         f"=ROUND(AVERAGE(Leads!{L['Rating']}2:{L['Rating']}{n}),2)", "0.00"),
        ("Resenas - mediana", f"=MEDIAN(Leads!{L['Resenas']}2:{L['Resenas']}{n})", "#,##0"),
        ("Resenas - total", f"=SUM(Leads!{L['Resenas']}2:{L['Resenas']}{n})", "#,##0"),
        ("Ciudades distintas",
         f'=SUMPRODUCT((Leads!{L["Ciudad"]}2:{L["Ciudad"]}{n}<>"")'
         f'/COUNTIF(Leads!{L["Ciudad"]}2:{L["Ciudad"]}{n},'
         f'Leads!{L["Ciudad"]}2:{L["Ciudad"]}{n}&""))', "#,##0"),
    ]
    r = 3
    for etiqueta, formula, fmt in metricas:
        rs.cell(r, 1, etiqueta).font = Font(name=FUENTE, size=11)
        b = rs.cell(r, 2, formula); b.font = Font(name=FUENTE, size=11, bold=True)
        if fmt:
            b.number_format = fmt
        r += 1

    r += 1
    for c_, txt in ((1, "Pipeline"), (2, "Leads"), (3, "% del total")):
        c = rs.cell(r, c_, txt)
        c.font = Font(name=FUENTE, bold=True, size=12, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    primera = r + 1
    for e in ESTADOS:
        r += 1
        rs.cell(r, 1, e).font = Font(name=FUENTE, size=11)
        rs.cell(r, 2, f"=COUNTIF(Leads!{letra_estado}2:{letra_estado}{n},A{r})"
                ).font = Font(name=FUENTE, size=11)
        p = rs.cell(r, 3, f"=IFERROR(B{r}/$B$3,0)")
        p.font = Font(name=FUENTE, size=11); p.number_format = "0.0%"
    r += 1
    rs.cell(r, 1, "Total").font = Font(name=FUENTE, bold=True, size=11)
    rs.cell(r, 2, f"=SUM(B{primera}:B{r-1})").font = Font(name=FUENTE, bold=True, size=11)
    tp = rs.cell(r, 3, f"=SUM(C{primera}:C{r-1})")
    tp.font = Font(name=FUENTE, bold=True, size=11); tp.number_format = "0.0%"

    # Desglose por zona buscada
    zonas = []
    for d in filas:
        if d["Zona buscada"] and d["Zona buscada"] not in zonas:
            zonas.append(d["Zona buscada"])
    if len(zonas) > 1:
        r += 2
        for c_, txt in ((1, "Zona buscada"), (2, "Leads")):
            c = rs.cell(r, c_, txt)
            c.font = Font(name=FUENTE, bold=True, size=12, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=AZUL)
        for z in zonas:
            r += 1
            rs.cell(r, 1, z).font = Font(name=FUENTE, size=11)
            rs.cell(r, 2,
                    f'=COUNTIF(Leads!{L["Zona buscada"]}2:{L["Zona buscada"]}{n},A{r})'
                    ).font = Font(name=FUENTE, size=11)

    # --- Fuente y metodo
    mt = wb.create_sheet("Fuente y metodo")
    mt.column_dimensions["A"].width = 26
    mt.column_dimensions["B"].width = 96
    mt.cell(1, 1, "Fuente, metodo y advertencias").font = Font(
        name=FUENTE, bold=True, size=14, color=GRIS)
    r = 3
    for k, v in meta:
        a = mt.cell(r, 1, k)
        a.font = Font(name=FUENTE, bold=True, size=10, color=GRIS)
        a.alignment = Alignment(vertical="top")
        b = mt.cell(r, 2, v)
        b.font = Font(name=FUENTE, size=10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if k.isupper():
            b.fill = PatternFill("solid", fgColor="FFF4CE")
        mt.row_dimensions[r].height = 15 + 15 * (len(v) // 110)
        r += 1

    wb.save(ruta)


def metadatos(args, country, runs, crudos, filas, desc):
    filtros = []
    if args.solo_con_telefono:
        filtros.append("solo con telefono normalizable")
    if args.min_rating:
        filtros.append(f"rating >= {args.min_rating}")
    if args.min_reviews:
        filtros.append(f"resenas >= {args.min_reviews}")
    if args.sin_web:
        filtros.append("solo negocios sin sitio web")
    filtros.append("cerrados omitidos" if args.skip_cerrados else "cerrados incluidos")
    detalle = ", ".join(f"{k}={v}" for k, v in desc.items() if v) or "ninguno"
    sin_wa = sum(1 for f in filas if not f["WhatsApp"])
    corridas = "; ".join(f"{z} -> {rid} ({n} lugares)" for z, rid, n, _ in runs)
    costo_real = sum(c for *_, c in runs)

    meta = [
        ("Fuente primaria",
         "Google Maps, extraido via Apify actor compass/crawler-google-places."),
        ("Corridas", corridas),
        ("Terminos de busqueda", " | ".join(args.niche)),
        ("Zonas", " | ".join(args.zone) + f" | countryCode={country.lower() or 'n/d'}"
                                          f" | language={args.language}"),
    ]
    if args.receta:
        meta.append(("Receta", f"{args.receta} - {RECETAS[args.receta]['nota']}"))
    meta += [
        ("Fecha de captura",
         f"{date.today().isoformat()} - los datos de Google Maps cambian; "
         f"revalidar telefono antes de contactar."),
        ("Volumen",
         f"Tope {args.max_results} por termino y zona; devueltos {len(crudos)}; "
         f"descartados: {detalle}; {len(filas)} en la hoja."),
        ("Filtros aplicados", "; ".join(filtros) + "."),
        ("Telefono",
         "Normalizado a E.164. El link de WhatsApp se construyo sobre ese numero, "
         "no sobre el texto crudo de Google. "
         + (f"{sin_wa} leads quedaron sin link." if sin_wa else
            "Todos los leads tienen link.")),
        ("Mensaje precargado", args.mensaje),
        ("COBERTURA",
         "Google Maps prioriza relevancia sobre exhaustividad: un tope bajo sesga hacia "
         "el area metropolitana. Para cobertura territorial real, agregar zonas "
         "(--zona A B C); la deduplicacion entre zonas ya es automatica."),
        ("CALIFICACION PENDIENTE",
         "El scraper devuelve categorias de Google, no capacidad instalada ni tamano "
         "real. Calificar antes de prospectar; fichas con muchas resenas suelen ser "
         "cadenas con una sola decision de compra."),
        ("CUMPLIMIENTO",
         "Telefonos de negocios unipersonales son datos personales (Costa Rica: Ley "
         "8968 y PRODHAB; UE: RGPD). El primer contacto debe identificar emisor y "
         "motivo y ofrecer baja explicita. Registrar y respetar toda solicitud de no "
         "contacto."),
        ("Costo de la corrida",
         (f"REAL facturado por Apify (usageTotalUsd): USD {costo_real:.4f} para "
          f"{len(crudos)} lugares = USD {costo_real / len(crudos):.4f} por lugar. "
          f"Tarifa de referencia del script: USD {USD_POR_LUGAR:.4f} por lugar "
          f"(medicion unica del 2026-07-28, n=1). "
          + ("Desvio >20%: actualizar USD_POR_LUGAR."
             if abs(costo_real / len(crudos) - USD_POR_LUGAR) / USD_POR_LUGAR > 0.20
             else "Sin desvio material.")
          if costo_real and crudos else
          f"Apify no reporto usageTotalUsd; estimado ~USD "
          f"{len(crudos) * USD_POR_LUGAR:.2f} ({len(crudos)} lugares x USD "
          f"{USD_POR_LUGAR:.4f}, medicion unica del 2026-07-28). "
          f"Verificar la tarifa vigente del actor.")),
        ("Respaldo crudo",
         "El archivo _raw.json conserva la respuesta integra de Apify. Guardalo: "
         "Apify borra los datasets a los 7 dias en plan FREE."),
    ]
    return meta


# ------------------------------------------------------------------ Main

def main():
    p = argparse.ArgumentParser(
        description="Buscador general de negocios en Google Maps -> XLSX + CSV")
    p.add_argument("--buscar", "--niche", dest="niche", nargs="+", default=None,
                   help='Uno o varios terminos. Ej: --buscar "clinica dental" "odontologo"')
    p.add_argument("--zona", "--zone", dest="zone", nargs="+", default=None,
                   help='Una o varias zonas. Ej: --zona "San Jose, CR" "Cartago, CR"')
    p.add_argument("--receta", default=None,
                   help="Paquete de terminos predefinido: " + ", ".join(RECETAS))
    p.add_argument("--listar-recetas", action="store_true")
    p.add_argument("--apify-key",
                   default=os.getenv("APIFY_TOKEN") or APIFY_TOKEN_EMBEBIDO)
    p.add_argument("--country", default=None, help="ISO2, ej CR. Evita geocodificar")
    p.add_argument("--max-results", type=int, default=200,
                   help="Tope POR termino y POR zona (default 200)")
    p.add_argument("--language", default="es")
    p.add_argument("--min-rating", type=float, default=None)
    p.add_argument("--min-reviews", type=int, default=None)
    p.add_argument("--sin-web", action="store_true")
    p.add_argument("--skip-cerrados", action="store_true", default=True)
    p.add_argument("--incluir-cerrados", dest="skip_cerrados", action="store_false")
    p.add_argument("--solo-con-telefono", action="store_true")
    p.add_argument("--mensaje", default="Hola, buenos dias. Le escribo de parte de "
                                        "nuestra empresa, hablo con {empresa}?")
    p.add_argument("--formato", choices=["xlsx", "csv", "ambos"], default="ambos")
    p.add_argument("--out-dir", default=".")
    p.add_argument("--nombre", default=None, help="Nombre base de los archivos")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--ignorar-presupuesto", action="store_true")
    args = p.parse_args()

    if args.listar_recetas:
        for k, v in RECETAS.items():
            print(f"\n{k}\n  terminos: {', '.join(v['terminos'])}\n  nota: {v['nota']}")
        return

    args = intake(args)
    if not args.apify_key:
        sys.exit("[X] Falta APIFY_TOKEN (env) o --apify-key")
    os.makedirs(args.out_dir, exist_ok=True)

    country = resolver_pais(args.zone[0], args.country)
    crudos, runs = scrapear(args.apify_key, args, country)
    if not crudos:
        sys.exit("[X] Cero resultados. Proba terminos mas genericos o zonas mas amplias.")

    print("[2/4] Normalizando y deduplicando (global entre terminos y zonas)...")
    filas, desc = construir(crudos, country, args)
    detalle = ", ".join(f"{k}={v}" for k, v in desc.items() if v)
    print(f"      {len(filas)} leads utiles" + (f" | descartados: {detalle}" if detalle else ""))
    sin_wa = sum(1 for f in filas if not f["WhatsApp"])
    if sin_wa:
        print(f"      [!] {sin_wa} sin WhatsApp (telefono ausente o no normalizable)")
    if not filas:
        sys.exit("[X] Ningun lead paso los filtros.")

    etiqueta = args.nombre or args.receta or args.niche[0]
    if not args.nombre:
        etiqueta = f"{etiqueta} {args.zone[0]}" + (" y mas" if len(args.zone) > 1 else "")
    slug = re.sub(r"[^a-z0-9]+", "-",
                  etiqueta.lower().translate(str.maketrans("áéíóúñ", "aeioun"))).strip("-")
    base = os.path.join(args.out_dir, f"leads-{slug}")

    print("[3/4] Respaldo crudo...")
    with open(f"{base}_raw.json", "w", encoding="utf-8") as fh:
        json.dump(crudos, fh, ensure_ascii=False, indent=1)
    print(f"      {base}_raw.json")

    print("[4/4] Entregables...")
    generados = [f"{base}_raw.json"]
    if args.formato in ("csv", "ambos"):
        escribir_csv(filas, f"{base}.csv"); generados.append(f"{base}.csv")
        print(f"      {base}.csv")
    if args.formato in ("xlsx", "ambos"):
        escribir_xlsx(filas, f"{base}.xlsx",
                      metadatos(args, country, runs, crudos, filas, desc))
        generados.append(f"{base}.xlsx")
        print(f"      {base}.xlsx  (Leads / Resumen / Fuente y metodo)")

    print(f"\n[OK] {len(filas)} leads. Archivos:")
    for g in generados:
        print("     " + g)
    print("     Las formulas del XLSX no traen valor en cache hasta que Excel o "
          "LibreOffice abran el archivo.")


if __name__ == "__main__":
    main()
